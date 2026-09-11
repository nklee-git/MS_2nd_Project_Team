
"""
blackup.kr 카테고리(cate_no=860) 상품 수집 스크립트
- 목록 페이지를 순회하며 상품 상세페이지 링크(product_no)를 수집
- 각 상세페이지의 JSON-LD(schema.org Product) 데이터를 파싱해
  상품명 / 대표이미지 / 가격 / 색상 / 사이즈 조합을 추출
- 결과를 엑셀(xlsx)로 저장
"""

import json
import re
import sys
import time

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

try:
    sys.stdout.reconfigure(encoding="utf-8")
except AttributeError:
    pass

BASE = "https://blackup.kr"
CATE_NO = 860
LIST_URL = f"{BASE}/product/list.html"
DETAIL_URL = f"{BASE}/product/detail.html"

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0 Safari/537.36"
    )
}

session = requests.Session()
session.headers.update(HEADERS)


def get_product_numbers(max_pages=30):
    """카테고리 목록 페이지를 페이지 번호를 늘려가며 순회, product_no 수집

    사이드바 추천/베스트/예약판매 섹션에도 다른 상품 링크가 섞여 있으므로,
    반드시 메인 상품 그리드(xans-product-listnormal) 안의 링크만 추출한다.
    """
    product_nos = []
    seen = set()
    page = 1
    while page <= max_pages:
        resp = session.get(
            LIST_URL,
            params={"cate_no": CATE_NO, "page": page},
            timeout=15,
        )
        resp.raise_for_status()
        soup = BeautifulSoup(resp.text, "html.parser")
        grid = soup.find("div", class_="xans-product-listnormal")
        if grid is None:
            break

        nos_in_page = []
        for a in grid.select("a[href*='product_no=']"):
            m = re.search(r"product_no=(\d+)", a.get("href", ""))
            if m and m.group(1) not in nos_in_page:
                nos_in_page.append(m.group(1))

        new_nos = [n for n in nos_in_page if n not in seen]
        if not new_nos:
            break
        for n in new_nos:
            seen.add(n)
            product_nos.append(n)
        print(f"[목록] {page}페이지 -> 신규 {len(new_nos)}개 (누적 {len(product_nos)}개)")
        page += 1
        time.sleep(0.3)
    return product_nos


def get_product_detail(product_no):
    """상세페이지에서 JSON-LD Product 데이터를 파싱"""
    resp = session.get(
        DETAIL_URL,
        params={"product_no": product_no, "cate_no": CATE_NO},
        timeout=15,
    )
    resp.raise_for_status()
    soup = BeautifulSoup(resp.text, "html.parser")

    ld_json = None
    for tag in soup.select('script[type="application/ld+json"]'):
        try:
            data = json.loads(tag.string or tag.text)
        except (json.JSONDecodeError, TypeError):
            continue
        if data.get("@type") == "Product":
            ld_json = data
            break

    if not ld_json:
        return None

    name = ld_json.get("name", "").strip()
    images = ld_json.get("image", [])
    main_image = images[0] if images else ""
    offers = ld_json.get("offers", [])

    if isinstance(offers, dict):
        offers = [offers]

    variants = []
    colors, sizes = [], []
    for offer in offers:
        if not isinstance(offer, dict):
            continue
        offer_name = offer.get("name", "")
        # "{상품명} {색상}-{사이즈}" 형태에서 색상/사이즈만 분리
        suffix = offer_name.replace(name, "", 1).strip()
        if "-" in suffix:
            color, size = suffix.rsplit("-", 1)
            color, size = color.strip(), size.strip()
        else:
            color, size = suffix, ""
        if color and color not in colors:
            colors.append(color)
        if size and size not in sizes:
            sizes.append(size)
        variants.append(
            {
                "color": color,
                "size": size,
                "price": offer.get("price"),
                "availability": offer.get("availability", "").replace(
                    "https://schema.org/", ""
                ),
                "url": offer.get("url", ""),
            }
        )

    base_price = min((v["price"] for v in variants if v["price"]), default=None)

    return {
        "product_no": product_no,
        "name": name,
        "image": main_image,
        "price": base_price,
        "colors": colors,
        "sizes": sizes,
        "variants": variants,
        "detail_url": f"{DETAIL_URL}?product_no={product_no}&cate_no={CATE_NO}",
    }


def main():
    product_nos = get_product_numbers()
    print(f"총 {len(product_nos)}개 상품 발견. 상세 정보 수집 시작...")

    products = []
    for i, no in enumerate(product_nos, 1):
        try:
            detail = get_product_detail(no)
        except requests.RequestException as e:
            print(f"  ! {no} 요청 실패: {e}")
            continue
        if detail:
            products.append(detail)
            print(f"  [{i}/{len(product_nos)}] {detail['name']}")
        else:
            print(f"  [{i}/{len(product_nos)}] {no} - JSON-LD 없음, 건너뜀")
        time.sleep(0.3)

    # ---- 엑셀 저장 (상품 요약 시트 + 옵션별 상세 시트) ----
    wb = Workbook()

    ws1 = wb.active
    ws1.title = "상품목록"
    ws1.append(["상품번호", "상품명", "가격", "색상", "사이즈", "대표이미지", "상품링크"])
    for p in products:
        ws1.append(
            [
                p["product_no"],
                p["name"],
                p["price"],
                ", ".join(p["colors"]),
                ", ".join(p["sizes"]),
                p["image"],
                p["detail_url"],
            ]
        )

    ws2 = wb.create_sheet("옵션상세")
    ws2.append(["상품번호", "상품명", "색상", "사이즈", "가격", "재고상태", "옵션링크"])
    for p in products:
        for v in p["variants"]:
            ws2.append(
                [
                    p["product_no"],
                    p["name"],
                    v["color"],
                    v["size"],
                    v["price"],
                    v["availability"],
                    v["url"],
                ]
            )

    for ws in (ws1, ws2):
        for col_idx, col_cells in enumerate(ws.columns, 1):
            max_len = max((len(str(c.value)) if c.value is not None else 0) for c in col_cells)
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 60)

    out_path = "blackup_products.xlsx"
    wb.save(out_path)
    print(f"\n완료: {len(products)}개 상품 -> {out_path}")


if __name__ == "__main__":
    main()
